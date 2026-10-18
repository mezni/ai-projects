import os
import tempfile
import time

USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
os.environ.setdefault("USER_AGENT", USER_AGENT)

from langchain_classic.chains import (
    create_history_aware_retriever,
    create_retrieval_chain,
)
from langchain_classic.chains.combine_documents import create_stuff_documents_chain
from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    WebBaseLoader,
)
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import settings
from llm_client import get_embeddings, get_llm

BROWSER_HEADERS = {"User-Agent": USER_AGENT}


def load_docx(file_path):
    """Extracts text from a .docx using python-docx (no unstructured needed)."""
    import docx

    doc = docx.Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("No extractable text found in the .docx file.")
    return [Document(page_content=text, metadata={"source": file_path})]


def load_source_documents(uploaded_file, url_input, source_type):
    """Safely extracts documents from uploaded files or web links."""
    documents = []
    try:
        if source_type == "File Upload" and uploaded_file:
            file_ext = os.path.splitext(uploaded_file.name)[1].lower()
            with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp_file:
                tmp_file.write(uploaded_file.getvalue())
                tmp_path = tmp_file.name

            try:
                if file_ext == ".pdf":
                    documents = PyPDFLoader(tmp_path).load()
                elif file_ext == ".txt":
                    documents = TextLoader(tmp_path, encoding="utf-8").load()
                elif file_ext == ".docx":
                    documents = load_docx(tmp_path)
                else:
                    raise ValueError(f"Unsupported file format: {file_ext}")
            finally:
                if os.path.exists(tmp_path):
                    os.remove(tmp_path)

        elif source_type == "HTTP URL" and url_input.strip():
            url = url_input.strip()
            if not url.startswith(("http://", "https://")):
                url = f"https://{url}"
            loader = WebBaseLoader(
                url,
                header_template=BROWSER_HEADERS,
                requests_kwargs={"timeout": 30},
            )
            documents = loader.load()

        if not documents or not any(d.page_content.strip() for d in documents):
            raise ValueError("No readable content could be extracted from the source.")

    except Exception as e:
        raise RuntimeError(f"Error loading source data: {e}") from e

    return documents


def initialize_vectorstore(documents):
    """Chunks documents, builds a FAISS vector store, and returns chunk metrics."""
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE, chunk_overlap=settings.CHUNK_OVERLAP
    )

    chunk_start = time.perf_counter()
    splits = text_splitter.split_documents(documents)
    chunking_time = time.perf_counter() - chunk_start

    if not splits:
        raise ValueError("No text chunks could be created from the source.")

    index_start = time.perf_counter()
    embeddings = get_embeddings()
    vectorstore = FAISS.from_documents(splits, embeddings)
    indexing_time = time.perf_counter() - index_start

    chunk_info = {
        "total_chunks": len(splits),
        "chunk_size": settings.CHUNK_SIZE,
        "chunk_overlap": settings.CHUNK_OVERLAP,
        "total_docs": len(documents),
        "total_chars": sum(len(d.page_content) for d in documents),
        "avg_chunk_chars": round(
            sum(len(s.page_content) for s in splits) / len(splits)
        ),
        "chunking_time": round(chunking_time, 2),
        "indexing_time": round(indexing_time, 2),
    }

    return vectorstore, chunk_info


def build_rag_chain(vectorstore):
    """Constructs a conversational history-aware RAG pipeline."""
    llm = get_llm(temperature=0.2, streaming=True)
    retriever = vectorstore.as_retriever(search_kwargs={"k": 4})

    contextualize_q_prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                (
                    "Given a chat history and the latest user question which might reference context in the chat history, "
                    "formulate a standalone question which can be understood without the chat history. Do NOT answer, just reformulate if needed."
                ),
            ),
            MessagesPlaceholder("chat_history"),
            ("human", "{input}"),
        ]
    )
    history_aware_retriever = create_history_aware_retriever(
        llm, retriever, contextualize_q_prompt
    )

    qa_prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                (
                    "You are a helpful assistant. Use the following retrieved context to answer the question concisely (max 3 sentences). "
                    "If you don't know, say you don't know.\n\nContext:\n{context}"
                ),
            ),
            MessagesPlaceholder("chat_history"),
            ("human", "{input}"),
        ]
    )
    question_answer_chain = create_stuff_documents_chain(llm, qa_prompt)

    return create_retrieval_chain(history_aware_retriever, question_answer_chain)
